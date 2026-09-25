from datetime import date

import pytest

from packages.contracts.canonical import BlockKind
from packages.contracts.source import SourceDocument, SourceFormat
from services.ingestion.extractors.fixtures import FixtureDocumentParser
from services.ingestion.structure.canonical_document import Canonicalizer
from services.ingestion.structure.markdown import canonical_to_markdown


def markdown_source() -> SourceDocument:
    return SourceDocument(
        id="source-markdown",
        organization_id="ajt",
        policy_id="policy-markdown",
        version_id="version-markdown",
        file_name="verified-policy.md",
        media_type="text/markdown",
        source_format=SourceFormat.MARKDOWN,
        sha256="markdown",
        original_artifact_uri="local://ajt/verified-policy.md",
    )


async def test_markdown_preserves_heading_list_and_source_anchor() -> None:
    source = SourceDocument(
        id="source-1",
        organization_id="ajt",
        policy_id="policy-1",
        version_id="version-1",
        file_name="fixture.md",
        media_type="text/markdown",
        source_format=SourceFormat.MARKDOWN,
        sha256="abc",
        original_artifact_uri="local://ajt/source",
    )
    raw = await FixtureDocumentParser().parse(
        source,
        b"# Store Operations\n## 4.3 Damaged Stock\n1. Isolate the item\n2. Record damage",
    )
    canonical = Canonicalizer().canonicalize(source, raw)

    damaged = next(section for section in canonical.sections if "Damaged" in section.heading)
    assert damaged.heading_level == 2
    assert damaged.heading_path == ("Store Operations", "4.3 Damaged Stock")
    assert damaged.policy_number == "4.3"
    assert damaged.blocks[0].list_items[0].source.block_anchor == "line-3"


@pytest.mark.parametrize(
    ("heading", "expected"),
    [
        ("# SOP #76 Store Safety", "SOP #76"),
        ("# SOP-76 Store Safety", "SOP-76"),
        ("# Policy #93 Leave", "Policy #93"),
    ],
)
async def test_markdown_compatibility_policy_number_variants(
    heading: str, expected: str
) -> None:
    raw = await FixtureDocumentParser().parse(
        markdown_source(), f"{heading}\nPolicy content".encode()
    )

    assert raw.policy_number == expected


async def test_verified_markdown_metadata_hierarchy_lists_and_table() -> None:
    content = b"""---
title: Store Leave Policy
policy_number: SOP-76
effective_date: 2026-09-15
---
# Store Leave Policy
Verified policy introduction.
## 1. Eligibility
1. Submit a request.
  1. Add the required details.
2. Wait for approval.
- Permanent employee
  - Confirmed status
### 1.1 Approval Matrix
| Role | Maximum days |
| --- | --- |
| Manager | 5 |
| Director | 10 |
"""
    raw = await FixtureDocumentParser().parse(markdown_source(), content)
    canonical = Canonicalizer().canonicalize(markdown_source(), raw)

    assert raw.title == "Store Leave Policy"
    assert raw.policy_number == "SOP-76"
    assert raw.effective_date == date(2026, 9, 15)
    assert canonical.title == "Store Leave Policy"
    assert canonical.policy_number == "SOP-76"
    assert canonical.effective_date == date(2026, 9, 15)

    eligibility = next(
        section for section in canonical.sections if section.heading.endswith("Eligibility")
    )
    assert eligibility.policy_number == "1"
    assert eligibility.heading_path == ("Store Leave Policy", "1. Eligibility")
    ordered = next(block for block in eligibility.blocks if block.kind is BlockKind.ORDERED_LIST)
    unordered = next(
        block for block in eligibility.blocks if block.kind is BlockKind.UNORDERED_LIST
    )
    assert [item.level for item in ordered.list_items] == [0, 1, 0]
    assert [item.level for item in unordered.list_items] == [0, 1]

    matrix = next(section for section in canonical.sections if "Approval Matrix" in section.heading)
    assert matrix.policy_number == "1.1"
    assert matrix.heading_level == 3
    assert matrix.heading_path == (
        "Store Leave Policy",
        "1. Eligibility",
        "1.1 Approval Matrix",
    )
    table_block = next(block for block in matrix.blocks if block.kind is BlockKind.TABLE)
    assert table_block.table is not None
    assert [(cell.text, cell.is_header) for cell in table_block.table.cells] == [
        ("Role", True),
        ("Maximum days", True),
        ("Manager", False),
        ("5", False),
        ("Director", False),
        ("10", False),
    ]
    assert table_block.source.block_anchor == "lines-15-18"
    rendered = canonical_to_markdown(canonical)
    assert "| Role | Maximum days |" in rendered
    assert "| Manager | 5 |" in rendered


async def test_markdown_missing_metadata_stays_null_and_uses_h1_title() -> None:
    raw = await FixtureDocumentParser().parse(
        markdown_source(), b"# General Guidance\n## 1. Eligibility\nPolicy text"
    )
    canonical = Canonicalizer().canonicalize(markdown_source(), raw)

    assert raw.title == "General Guidance"
    assert raw.policy_number is None
    assert raw.effective_date is None
    assert canonical.policy_number is None
    assert canonical.effective_date is None


async def test_xlsx_preserves_sheet_and_cell_coordinates() -> None:
    import io

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Approval Limits"
    sheet.append(["Role", "Limit"])
    sheet.append(["Manager", 5000])
    content = io.BytesIO()
    workbook.save(content)
    source = SourceDocument(
        id="source-2",
        organization_id="ajt",
        policy_id="policy-1",
        version_id="version-1",
        file_name="limits.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        source_format=SourceFormat.XLSX,
        sha256="def",
        original_artifact_uri="local://ajt/source",
    )
    raw = await FixtureDocumentParser().parse(source, content.getvalue())
    canonical = Canonicalizer().canonicalize(source, raw)
    table = canonical.sections[0].blocks[0].table
    assert table is not None
    cell = table.cells[0]

    assert cell.source.sheet_name == "Approval Limits"
    assert cell.source.cell_range == "A1"
