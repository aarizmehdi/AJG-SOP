import io
import re
from dataclasses import dataclass
from datetime import date, datetime

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from packages.contracts.source import ParserCapability, SourceDocument, SourceFormat
from services.ingestion.extractors.base import (
    DocumentParser,
    ParserUnavailableError,
    RawBlock,
    RawBlockKind,
    RawDocumentResult,
    RawTableCell,
)


@dataclass(frozen=True)
class MarkdownParseResult:
    blocks: list[RawBlock]
    warnings: list[str]
    title: str
    policy_number: str | None
    effective_date: date | None


class FixtureDocumentParser(DocumentParser):
    """Deterministic local parser for architecture tests; quality is not production-validated."""

    _supported = {
        SourceFormat.PDF,
        SourceFormat.DOCX,
        SourceFormat.XLSX,
        SourceFormat.MARKDOWN,
        SourceFormat.STRUCTURED_TEXT,
    }

    def capabilities(self, organization_id: str) -> list[ParserCapability]:
        return [
            ParserCapability(
                organization_id=organization_id,
                provider="fixture",
                source_format=source_format,
                available=source_format in self._supported,
                detail=(
                    "Deterministic local extraction; requires human review"
                    if source_format in self._supported
                    else "Requires a configured OCR-capable provider"
                ),
            )
            for source_format in SourceFormat
        ]

    async def parse(self, source: SourceDocument, content: bytes) -> RawDocumentResult:
        fallback_title = source.file_name.rsplit(".", 1)[0]
        if source.source_format is SourceFormat.MARKDOWN:
            parsed = self._parse_markdown(content, fallback_title)
            return RawDocumentResult(
                organization_id=source.organization_id,
                provider="fixture",
                provider_version="2",
                source_document_id=source.id,
                title=parsed.title,
                policy_number=parsed.policy_number,
                effective_date=parsed.effective_date,
                blocks=parsed.blocks,
                warnings=parsed.warnings,
            )
        parsers = {
            SourceFormat.PDF: self._parse_pdf,
            SourceFormat.DOCX: self._parse_docx,
            SourceFormat.XLSX: self._parse_xlsx,
            SourceFormat.STRUCTURED_TEXT: self._parse_text,
        }
        parser = parsers.get(source.source_format)
        if parser is None:
            raise ParserUnavailableError(
                f"Fixture parser cannot OCR {source.source_format.value}; "
                "configure Azure or Docling"
            )
        blocks, warnings = parser(content)
        return RawDocumentResult(
            organization_id=source.organization_id,
            provider="fixture",
            provider_version="1",
            source_document_id=source.id,
            title=fallback_title,
            blocks=blocks,
            warnings=warnings,
        )

    @classmethod
    def _parse_markdown(cls, content: bytes, fallback_title: str) -> MarkdownParseResult:
        lines = content.decode("utf-8-sig").splitlines()
        metadata, content_start, warnings = cls._front_matter(lines)
        title = metadata.get("title", "").strip() or None
        policy_number = (
            metadata.get("policy_number")
            or metadata.get("sop_number")
            or metadata.get("policy")
        )
        policy_number = policy_number.strip() if policy_number else None
        date_value = metadata.get("effective_date") or metadata.get("issue_date")
        effective_date = cls._parse_date(date_value) if date_value else None
        if date_value and effective_date is None:
            warnings.append("Markdown metadata date could not be parsed; effective_date is null")

        blocks: list[RawBlock] = []
        index = content_start
        while index < len(lines):
            line = lines[index]
            stripped = line.strip()
            if not stripped:
                index += 1
                continue

            heading = re.match(r"^(#{1,6})\s+(.+?)\s*#*$", stripped)
            if heading:
                heading_text = heading.group(2).strip()
                blocks.append(
                    RawBlock(
                        kind=RawBlockKind.HEADING,
                        text=heading_text,
                        level=len(heading.group(1)),
                        block_anchor=f"line-{index + 1}",
                    )
                )
                if title is None and len(heading.group(1)) == 1:
                    title = cls._title_without_policy_number(heading_text)
                if policy_number is None:
                    policy_number = cls._labelled_policy_number(heading_text)
                index += 1
                continue

            if index + 1 < len(lines) and cls._is_table(lines[index], lines[index + 1]):
                table, next_index = cls._parse_table(lines, index)
                blocks.append(table)
                index = next_index
                continue

            list_item = re.match(r"^((?:\d+[.)])|[-*+])\s+(.+)$", stripped)
            if list_item:
                blocks.append(
                    RawBlock(
                        kind=RawBlockKind.LIST_ITEM,
                        text=list_item.group(2),
                        marker=list_item.group(1),
                        level=max(0, (len(line) - len(line.lstrip())) // 2),
                        block_anchor=f"line-{index + 1}",
                    )
                )
                index += 1
                continue

            if policy_number is None:
                policy_number = cls._labelled_policy_number(stripped)
            if effective_date is None:
                inline_date = re.match(
                    r"^(?:effective|issue)\s+date\s*:\s*(.+)$", stripped, re.IGNORECASE
                )
                if inline_date:
                    effective_date = cls._parse_date(inline_date.group(1))
                    if effective_date is None:
                        warnings.append(
                            f"Line {index + 1} date could not be parsed; effective_date is null"
                        )
            blocks.append(
                RawBlock(
                    kind=RawBlockKind.PARAGRAPH,
                    text=stripped,
                    block_anchor=f"line-{index + 1}",
                )
            )
            index += 1

        return MarkdownParseResult(
            blocks=blocks,
            warnings=warnings,
            title=title or fallback_title,
            policy_number=policy_number,
            effective_date=effective_date,
        )

    @staticmethod
    def _front_matter(lines: list[str]) -> tuple[dict[str, str], int, list[str]]:
        if not lines or lines[0].strip() != "---":
            return {}, 0, []
        metadata: dict[str, str] = {}
        warnings: list[str] = []
        for index in range(1, min(len(lines), 51)):
            stripped = lines[index].strip()
            if stripped == "---":
                return metadata, index + 1, warnings
            if not stripped or stripped.startswith("#"):
                continue
            match = re.match(r"^([A-Za-z_][A-Za-z0-9_-]*)\s*:\s*(.*)$", stripped)
            if not match:
                warnings.append(f"Markdown metadata line {index + 1} was ignored")
                continue
            key = match.group(1).casefold().replace("-", "_")
            value = match.group(2).strip().strip("\"'")
            allowed_keys = {
                "title",
                "policy_number",
                "sop_number",
                "policy",
                "effective_date",
                "issue_date",
            }
            if key in allowed_keys:
                metadata[key] = value
        return {}, 0, ["Markdown front matter has no closing delimiter"]

    @staticmethod
    def _parse_date(value: str) -> date | None:
        cleaned = value.strip()
        for format_string in ("%Y-%m-%d", "%d %B %Y", "%d %b %Y", "%B %d, %Y", "%b %d, %Y"):
            try:
                return datetime.strptime(cleaned, format_string).date()
            except ValueError:
                continue
        return None

    @staticmethod
    def _labelled_policy_number(value: str) -> str | None:
        match = re.search(
            r"\b(?:SOP|Policy)\s*(?:No\.?|Number|#|-)?\s*"
            r"(?=[A-Za-z0-9./-]*\d)[A-Za-z0-9][A-Za-z0-9./-]*",
            value,
            flags=re.IGNORECASE,
        )
        return re.sub(r"\s+", " ", match.group(0).strip()) if match else None

    @classmethod
    def _title_without_policy_number(cls, value: str) -> str:
        policy_number = cls._labelled_policy_number(value)
        if not policy_number:
            return value.strip()
        title = value.replace(policy_number, "", 1).strip(" :-–—")
        return title or value.strip()

    @classmethod
    def _is_table(cls, header: str, delimiter: str) -> bool:
        headers = cls._split_table_row(header)
        delimiters = cls._split_table_row(delimiter)
        return bool(
            "|" in header
            and headers
            and len(headers) == len(delimiters)
            and all(re.fullmatch(r":?-{3,}:?", item) for item in delimiters)
        )

    @classmethod
    def _parse_table(cls, lines: list[str], start: int) -> tuple[RawBlock, int]:
        rows = [cls._split_table_row(lines[start])]
        index = start + 2
        while index < len(lines):
            stripped = lines[index].strip()
            if not stripped or "|" not in stripped:
                break
            rows.append(cls._split_table_row(lines[index]))
            index += 1
        column_count = max(len(row) for row in rows)
        cells = [
            RawTableCell(
                row=row_index,
                column=column_index,
                text=row[column_index] if column_index < len(row) else "",
                is_header=row_index == 0,
                cell_reference=(
                    f"L{start + 1 + (0 if row_index == 0 else row_index + 1)}"
                    f"C{column_index + 1}"
                ),
            )
            for row_index, row in enumerate(rows)
            for column_index in range(column_count)
        ]
        return (
            RawBlock(
                kind=RawBlockKind.TABLE,
                table_cells=cells,
                block_anchor=f"lines-{start + 1}-{index}",
            ),
            index,
        )

    @staticmethod
    def _split_table_row(line: str) -> list[str]:
        stripped = line.strip()
        if stripped.startswith("|"):
            stripped = stripped[1:]
        if stripped.endswith("|") and not stripped.endswith(r"\|"):
            stripped = stripped[:-1]
        parts = re.split(r"(?<!\\)\|", stripped)
        return [part.replace(r"\|", "|").strip() for part in parts]

    @staticmethod
    def _parse_text(content: bytes) -> tuple[list[RawBlock], list[str]]:
        blocks: list[RawBlock] = []
        for index, line in enumerate(content.decode("utf-8").splitlines()):
            stripped = line.strip()
            if not stripped:
                continue
            heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
            list_item = re.match(r"^((?:\d+[.)])|[-*])\s+(.+)$", stripped)
            if heading:
                blocks.append(
                    RawBlock(
                        kind=RawBlockKind.HEADING,
                        text=heading.group(2),
                        level=len(heading.group(1)),
                        block_anchor=f"line-{index + 1}",
                    )
                )
            elif list_item:
                marker = list_item.group(1)
                blocks.append(
                    RawBlock(
                        kind=RawBlockKind.LIST_ITEM,
                        text=list_item.group(2),
                        marker=marker,
                        level=max(0, (len(line) - len(line.lstrip())) // 2),
                        block_anchor=f"line-{index + 1}",
                    )
                )
            else:
                blocks.append(
                    RawBlock(
                        kind=RawBlockKind.PARAGRAPH,
                        text=stripped,
                        block_anchor=f"line-{index + 1}",
                    )
                )
        return blocks, []

    @staticmethod
    def _parse_pdf(content: bytes) -> tuple[list[RawBlock], list[str]]:
        reader = PdfReader(io.BytesIO(content))
        blocks: list[RawBlock] = []
        warnings: list[str] = []
        for page_index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if not text:
                warnings.append(f"Page {page_index} contains no native text and requires OCR")
                continue
            blocks.extend(
                RawBlock(kind=RawBlockKind.PARAGRAPH, text=line, page=page_index)
                for line in text.splitlines()
                if line.strip()
            )
        return blocks, warnings

    @staticmethod
    def _parse_docx(content: bytes) -> tuple[list[RawBlock], list[str]]:
        document = Document(io.BytesIO(content))
        blocks: list[RawBlock] = []
        for index, paragraph in enumerate(document.paragraphs):
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name if paragraph.style else ""
            match = re.match(r"Heading\s+(\d+)", style)
            kind = RawBlockKind.HEADING if match else RawBlockKind.PARAGRAPH
            blocks.append(
                RawBlock(
                    kind=kind,
                    text=text,
                    level=int(match.group(1)) if match else None,
                    block_anchor=f"paragraph-{index + 1}",
                )
            )
        for table_index, table in enumerate(document.tables):
            cells = [
                RawTableCell(
                    row=row_index,
                    column=column_index,
                    text=cell.text.strip(),
                    is_header=row_index == 0,
                )
                for row_index, row in enumerate(table.rows)
                for column_index, cell in enumerate(row.cells)
            ]
            blocks.append(
                RawBlock(
                    kind=RawBlockKind.TABLE,
                    table_cells=cells,
                    block_anchor=f"table-{table_index + 1}",
                )
            )
        return blocks, []

    @staticmethod
    def _parse_xlsx(content: bytes) -> tuple[list[RawBlock], list[str]]:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        blocks: list[RawBlock] = []
        for sheet in workbook.worksheets:
            cells = [
                RawTableCell(
                    row=cell.row - 1,
                    column=cell.column - 1,
                    text=str(cell.value) if cell.value is not None else "",
                    is_header=cell.row == 1,
                    sheet_name=sheet.title,
                    cell_reference=cell.coordinate,
                )
                for row in sheet.iter_rows()
                for cell in row
                if cell.value is not None
            ]
            if cells:
                blocks.append(
                    RawBlock(
                        kind=RawBlockKind.TABLE,
                        table_cells=cells,
                        sheet_name=sheet.title,
                        cell_range=sheet.calculate_dimension(),
                    )
                )
        return blocks, []
